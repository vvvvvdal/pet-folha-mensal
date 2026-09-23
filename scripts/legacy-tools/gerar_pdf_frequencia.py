#!/usr/bin/env python3
"""
Gera o PDF oficial da Folha de Frequência Mensal do PET-Saúde Clima UFG
seguindo com 100% de fidelidade o layout e padrão do arquivo:
'PET - Folha de Frequência Mensal.docx'
"""

import os
import sys
import zipfile
import subprocess
import copy
import xml.etree.ElementTree as ET
import openpyxl

DOCX_TEMPLATE = 'PET - Folha de Frequência Mensal.docx'
XLSX_DATA = 'PET - Folha de Frequência Mensal.xlsx'
OUTPUT_DOCX = 'PET - Folha de Frequência Mensal - Setembro 2026.docx'
OUTPUT_PDF = 'PET - Folha de Frequência Mensal - Setembro 2026.pdf'

W_NS = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
NS = {'w': W_NS}

def set_cell_text(tc, text):
    """Substitui o texto de uma célula w:tc preservando estilo w:rPr e parágrafo"""
    p = tc.find('w:p', NS)
    if p is None:
        p = ET.SubElement(tc, f'{{{W_NS}}}p')
    
    # Procura run existente para herdar estilos
    r = p.find('w:r', NS)
    rPr = None
    if r is not None:
        rPr = r.find('w:rPr', NS)
    
    # Limpa runs existentes
    for child in list(p):
        if child.tag == f'{{{W_NS}}}r':
            p.remove(child)
            
    # Cria novo run com o texto
    new_r = ET.SubElement(p, f'{{{W_NS}}}r')
    if rPr is not None:
        new_r.append(copy.deepcopy(rPr))
    new_t = ET.SubElement(new_r, f'{{{W_NS}}}t')
    new_t.text = str(text) if text else ''

def calc_pet_hours(arrival, departure):
    """Calcula horas PET: cada hora do relógio iniciada conta como 1h inteira"""
    if arrival is None or departure is None:
        return 0
    if hasattr(arrival, 'hour'):
        h1, m1 = arrival.hour, arrival.minute
    else:
        parts = str(arrival).strip().split(':')
        h1, m1 = int(parts[0]), int(parts[1]) if len(parts) > 1 else 0
        
    if hasattr(departure, 'hour'):
        h2, m2 = departure.hour, departure.minute
    else:
        parts = str(departure).strip().split(':')
        h2, m2 = int(parts[0]), int(parts[1]) if len(parts) > 1 else 0
        
    start_hour = h1
    end_hour = h2 + 1 if m2 > 0 else h2
    diff = end_hour - start_hour
    if diff < 0:
        diff += 24
    return max(0, diff)

def get_activities_from_xlsx(xlsx_path):
    """Lê as atividades e total da aba Lançamentos da planilha Excel"""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws_lanc = wb['Lançamentos']
    
    activities = []
    total_hours = 0
    
    for r in range(9, 44):
        dt_val = ws_lanc.cell(r, 1).value
        ch_val = ws_lanc.cell(r, 2).value
        sd_val = ws_lanc.cell(r, 3).value
        at_val = ws_lanc.cell(r, 4).value
        hr_val = ws_lanc.cell(r, 6).value
        
        if dt_val is not None and ch_val is not None and sd_val is not None:
            dt_str = dt_val.strftime('%d/%m/%Y') if hasattr(dt_val, 'strftime') else str(dt_val)[:10]
            ch_str = ch_val.strftime('%H:%M') if hasattr(ch_val, 'strftime') else str(ch_val)[:5]
            sd_str = sd_val.strftime('%H:%M') if hasattr(sd_val, 'strftime') else str(sd_val)[:5]
            
            try:
                h = int(hr_val) if hr_val not in (None, "") else calc_pet_hours(ch_val, sd_val)
            except (ValueError, TypeError):
                h = calc_pet_hours(ch_val, sd_val)
                
            total_hours += h
            
            activities.append({
                'data': dt_str,
                'chegada': ch_str,
                'saida': f"{sd_str} ({h}h)",
                'atividade': str(at_val),
                'assinatura': ''
            })
            
    return activities, total_hours

def generate_docx_and_pdf():
    if not os.path.exists(DOCX_TEMPLATE):
        print(f"❌ Template {DOCX_TEMPLATE} não encontrado!")
        sys.exit(1)
        
    activities, total_hours = get_activities_from_xlsx(XLSX_DATA)
    print(f"📋 Carregadas {len(activities)} atividades. Total de horas: {total_hours}h.")
    
    # 1. Descompactar e modificar document.xml
    with zipfile.ZipFile(DOCX_TEMPLATE, 'r') as zin:
        xml_content = zin.read('word/document.xml')
        all_files = {name: zin.read(name) for name in zin.namelist()}
        
    root = ET.fromstring(xml_content)
    
    # Localizar a tabela
    table = root.find('.//w:tbl', NS)
    if table is None:
        print("❌ Tabela não encontrada no documento!")
        sys.exit(1)
        
    rows = table.findall('w:tr', NS)
    header_row = rows[0]
    template_row = rows[1] # row com estilos padrão das células
    
    # Remover linhas de dados antigas (índice 1 em diante)
    for r in rows[1:]:
        table.remove(r)
        
    # Adicionar as linhas preenchidas
    for act in activities:
        new_row = copy.deepcopy(template_row)
        cells = new_row.findall('w:tc', NS)
        if len(cells) >= 5:
            set_cell_text(cells[0], act['data'])
            set_cell_text(cells[1], act['chegada'])
            set_cell_text(cells[2], act['saida'])
            set_cell_text(cells[3], act['atividade'])
            set_cell_text(cells[4], act['assinatura'])
        table.append(new_row)
        
    # Atualizar linha do TOTAL
    for p in root.findall('.//w:p', NS):
        p_text = ''.join([e.text for e in p.iter() if e.text])
        if 'TOTAL:' in p_text:
            replaced = False
            for r in p.findall('w:r', NS):
                t = r.find('w:t', NS)
                if t is not None and ('hora' in t.text or '4' in t.text or '0' in t.text):
                    t.text = f"  {total_hours} horas"
                    replaced = True
            if not replaced:
                new_r = ET.SubElement(p, f'{{{W_NS}}}r')
                new_t = ET.SubElement(new_r, f'{{{W_NS}}}t')
                new_t.text = f"  {total_hours} horas"
                    
    # Salvar novo DOCX
    all_files['word/document.xml'] = ET.tostring(root, encoding='utf-8')
    
    with zipfile.ZipFile(OUTPUT_DOCX, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in all_files.items():
            zout.writestr(name, data)
            
    print(f"✅ Documento Word gerado: {OUTPUT_DOCX}")
    
    # 2. Converter para PDF com LibreOffice
    cmd = ['libreoffice', '--headless', '--convert-to', 'pdf', OUTPUT_DOCX]
    print("🚀 Convertendo para PDF via LibreOffice...")
    result = subprocess.run(cmd, stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True)
    
    if result.returncode == 0:
        print(f"🎉 PDF Oficial gerado com sucesso: {OUTPUT_PDF}")
    else:
        print(f"⚠️ Erro ao converter via LibreOffice: {result.stderr}")

if __name__ == '__main__':
    generate_docx_and_pdf()
